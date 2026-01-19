#!/usr/bin/env python3
"""
Extract motion-line information from PPTX files and compute two line metrics:
1) main line  - straight line from shape's connector endpoints (if available),
   otherwise fitted (via PCA) to ALL path points, otherwise bbox-based.
2) tail line  - straight line defined by the LAST TWO path points.

OUTPUT (CSV – Option A):
- One CSV file per PPTX.
"""

import argparse
from pathlib import Path
from typing import List, Tuple, Optional, Dict

import numpy as np
import pandas as pd
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

EMU_PER_IN = 914400.0
DPI = 96.0
EMU_PER_PX = EMU_PER_IN / DPI  # 9525.0


def emu_to_px(v: float) -> float:
    return float(v) / EMU_PER_PX


# ---------- Shape filtering ----------

def is_candidate_motion_line(shape) -> bool:
    st = shape.shape_type
    name = (getattr(shape, "name", "") or "").lower()

    if st == MSO_SHAPE_TYPE.LINE:
        return True

    keywords = (
        "connector", "arrow", "line", "freeform",
        "scribble", "curve", "zigzag", "path",
    )
    return any(k in name for k in keywords)


# ---------- Geometry extraction ----------

def extract_path_points_px(shape, max_points: Optional[int] = None) -> List[Tuple[float, float]]:
    shape_elm = shape.element
    try:
        path_elems = shape_elm.xpath(".//a:path")
    except TypeError:
        return []

    if not path_elems:
        return []

    path = path_elems[0]
    w_attr, h_attr = path.get("w"), path.get("h")
    if not w_attr or not h_attr:
        return []

    w, h = float(w_attr), float(h_attr)
    if w == 0 or h == 0:
        return []

    left, top = float(shape.left), float(shape.top)
    width, height = float(shape.width), float(shape.height)
    rotation_deg = float(shape.rotation or 0.0)

    cx, cy = left + width / 2.0, top + height / 2.0
    theta = np.deg2rad(rotation_deg)
    cos_t, sin_t = np.cos(theta), np.sin(theta)

    def local_to_slide(x_local, y_local):
        x_emu = left + (x_local / w) * width
        y_emu = top + (y_local / h) * height
        dx, dy = x_emu - cx, y_emu - cy
        xr = cx + cos_t * dx - sin_t * dy
        yr = cy + sin_t * dx + cos_t * dy
        return emu_to_px(xr), emu_to_px(yr)

    points = []
    for child in path:
        tag = child.tag.split("}")[-1]
        if tag in ("moveTo", "lnTo", "cubicBezTo", "quadBezTo", "arcTo"):
            try:
                pt_elems = child.xpath(".//a:pt")
            except TypeError:
                continue
            for pt in pt_elems:
                px, py = local_to_slide(float(pt.get("x")), float(pt.get("y")))
                points.append((px, py))

    return points[-max_points:] if max_points and len(points) > max_points else points


# ---------- Line fitting ----------

def fit_principal_line(points):
    if len(points) < 2:
        return None
    pts = np.array(points, dtype=float)
    center = pts.mean(axis=0)
    X = pts - center
    _, _, vt = np.linalg.svd(X, full_matrices=False)
    vx, vy = vt[0]
    t = max(np.ptp(pts[:, 0]), np.ptp(pts[:, 1]), 1.0)
    x1, y1 = center[0] - vx * t, center[1] - vy * t
    x2, y2 = center[0] + vx * t, center[1] + vy * t
    slope = (y2 - y1) / (x2 - x1) if abs(x2 - x1) > 1e-9 else float("inf")
    return x1, y1, x2, y2, slope


def tail_line_from_last_two(points):
    if len(points) < 2:
        return None
    x1, y1 = points[-2]
    x2, y2 = points[-1]
    slope = (y2 - y1) / (x2 - x1) if abs(x2 - x1) > 1e-9 else float("inf")
    return x1, y1, x2, y2, slope


def line_from_bbox_fallback(shape):
    left, top = float(shape.left), float(shape.top)
    width, height = float(shape.width), float(shape.height)
    cx, cy = left + width / 2.0, top + height / 2.0
    half_len = max(width, height) / 2.0
    return emu_to_px(cx - half_len), emu_to_px(cy), emu_to_px(cx + half_len), emu_to_px(cy)


def connector_endpoints_px(shape):
    if all(hasattr(shape, a) for a in ("begin_x", "begin_y", "end_x", "end_y")):
        try:
            return (
                emu_to_px(float(shape.begin_x)),
                emu_to_px(float(shape.begin_y)),
                emu_to_px(float(shape.end_x)),
                emu_to_px(float(shape.end_y)),
            )
        except Exception:
            pass
    return None


# ---------- Core ----------

def process_pptx_file(pptx_path: Path) -> pd.DataFrame:
    prs = Presentation(str(pptx_path))
    rows = []

    for slide_idx, slide in enumerate(prs.slides, start=1):
        for shape_idx, shape in enumerate(slide.shapes, start=1):
            if not is_candidate_motion_line(shape):
                continue

            connector = connector_endpoints_px(shape)
            path_points = extract_path_points_px(shape)
            points_available = len(path_points)

            if connector:
                x1, y1, x2, y2 = connector
                slope = (y2 - y1) / (x2 - x1) if abs(x2 - x1) > 1e-9 else float("inf")
                source = "connector"
            else:
                main = fit_principal_line(path_points)
                if main:
                    x1, y1, x2, y2, slope = main
                    source = "path"
                else:
                    x1, y1, x2, y2 = line_from_bbox_fallback(shape)
                    slope = (y2 - y1) / (x2 - x1) if abs(x2 - x1) > 1e-9 else float("inf")
                    source = "bbox"

            tail = tail_line_from_last_two(path_points)
            if tail:
                tx1, ty1, tx2, ty2, tslope = tail
                tail_used = 2
            else:
                tx1 = ty1 = tx2 = ty2 = float("nan")
                tslope = float("inf")
                tail_used = 0

            rows.append({
                "pptx_file": pptx_path.name,
                "slide_index": slide_idx,
                "shape_index": shape_idx,
                "shape_name": getattr(shape, "name", ""),
                "shape_type": str(shape.shape_type),

                "line_x1_px": x1, "line_y1_px": y1,
                "line_x2_px": x2, "line_y2_px": y2,
                "line_slope": slope,
                "line_source": source,

                "tail_x1_px": tx1, "tail_y1_px": ty1,
                "tail_x2_px": tx2, "tail_y2_px": ty2,
                "tail_slope": tslope,
                "tail_points_available": points_available,
                "tail_points_used": tail_used,
            })

    return pd.DataFrame(rows)


# ---------- main ----------

def main():
    ap = argparse.ArgumentParser(
        description="Extract motion-line metrics from PPTX files (CSV output, one per PPTX)."
    )
    ap.add_argument("input_path", help="PPTX file or directory with PPTX files")
    ap.add_argument("-o", "--output-dir", default="motion_lines_metrics_csv")
    args = ap.parse_args()

    input_path = Path(args.input_path).resolve()
    output_dir = Path(args.output_dir).resolve()
    output_dir.mkdir(parents=True, exist_ok=True)

    pptx_files = (
        [input_path]
        if input_path.is_file()
        else sorted(input_path.glob("*.pptx"))
    )

    if not pptx_files:
        raise SystemExit("No PPTX files found")

    for pptx in pptx_files:
        print(f"Processing {pptx.name}")
        df = process_pptx_file(pptx)
        if df.empty:
            print("  -> no motion lines found")
            continue
        out_csv = output_dir / f"{pptx.stem}.csv"
        df.sort_values(["slide_index", "shape_index"], inplace=True)
        df.to_csv(out_csv, index=False, encoding="utf-8")
        print(f"  -> wrote {out_csv}")

    print(f"\nDone. CSVs saved to: {output_dir}")


if __name__ == "__main__":
    main()

#############
# HOW TO RUN:
#
# INPUT:
#   - A single PPTX file, OR
#   - A directory containing multiple PPTX files
#
# OUTPUT:
#   - A directory containing CSV files
#   - Each PPTX produces one CSV file with motion-line metrics
#
# ------------------------------------------------------------
#
# 1) Run on a SINGLE PPTX file:
#
#   python3 extract_motion_lines.py \
#      /full/path/to/presentation.pptx \
#     --output-dir motion_lines_metrics_csv
#
# 2) Run on a FOLDER of PPTX files:
#
#   python3 extract_motion_lines.py \
#      /full/path/to/pptx_folder \
#     --output-dir motion_lines_metrics_csv
#
# ------------------------------------------------------------
#
#   - Each output CSV contains:
#       * pptx_file, slide_index, shape_index, shape_name, shape_type
#       * main line endpoints and slope (line_x1_px, line_y1_px, line_x2_px, line_y2_px, line_slope)
#       * line_source (connector / path / bbox)
#       * tail line endpoints and slope (tail_x1_px, tail_y1_px, tail_x2_px, tail_y2_px, tail_slope)
#       * tail_points_available, tail_points_used
#
#############
